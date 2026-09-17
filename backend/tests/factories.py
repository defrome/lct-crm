"""Helpers for building test fixtures: xlsx files and catalog entities."""

from __future__ import annotations

import gzip
import io
import zipfile
from typing import Any

from openpyxl import Workbook

CATALOG_HEADERS = [
    "Название ВУЗа",
    "Вендор",
    "ПО",
    "Номер договора",
    "Подписание лицензии",
    "Срок действия лицензии (год)",
    "Статус по передаче",
    "ФИО Менеджера",
    "Ответственные от ВУЗа",
    "Комментарий",
]


def make_xlsx(rows: list[list[Any]], headers: list[str] | None = None) -> bytes:
    """Build an in-memory .xlsx with the customer's column titles."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Каталог"
    sheet.append(headers or CATALOG_HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def catalog_row(
    university: str,
    vendor: str | None = "Астра",
    product: str | None = "Astra Linux",
    contract: str | None = "ДЛ-001",
    signed: Any = "17.02.2026",
    years: Any = 3,
    status: str | None = "Передано",
    manager: str | None = None,
    contacts: str | None = None,
    comment: str | None = None,
) -> list[Any]:
    return [
        university,
        vendor,
        product,
        contract,
        signed,
        years,
        status,
        manager,
        contacts,
        comment,
    ]


def make_xls(rows: list[list[Any]], headers: list[str] | None = None) -> bytes:
    """Build an in-memory legacy .xls so the xlrd branch of the parser is covered.

    `xlwt` is a test-only dependency: the application reads .xls, never writes it.
    """
    import datetime as dt

    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("Каталог")
    date_style = xlwt.easyxf(num_format_str="DD.MM.YYYY")
    for column, title in enumerate(headers or CATALOG_HEADERS):
        sheet.write(0, column, title)
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            if isinstance(value, dt.date):
                sheet.write(row_index, column, value, date_style)
            else:
                sheet.write(row_index, column, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- attachment fixtures ---------------------------------------------------
# Minimal but genuine containers: the detector reads signatures and, for the
# compound formats, real archive entries — so these exercise the same path a
# user's file would.


def make_png() -> bytes:
    # 1x1 transparent PNG.
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
        "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )


def make_jpeg() -> bytes:
    return b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00" + b"\x00" * 32 + b"\xff\xd9"


def make_pdf() -> bytes:
    return b"%PDF-1.7\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"


def make_zip() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "содержимое")
    return buffer.getvalue()


def make_gzip() -> bytes:
    buffer = io.BytesIO()
    with gzip.GzipFile(fileobj=buffer, mode="wb") as handle:
        handle.write(b"content")
    return buffer.getvalue()


def make_rar() -> bytes:
    # RAR5 signature; the detector only identifies the container.
    return b"Rar!\x1a\x07\x01\x00" + b"\x00" * 64


def make_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
    return buffer.getvalue()


def make_xlsx_sample() -> bytes:
    workbook = Workbook()
    workbook.active.append(["значение"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_xls_sample() -> bytes:
    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    book.add_sheet("Лист").write(0, 0, "значение")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_doc() -> bytes:
    """An OLE2 container carrying a `WordDocument` stream name.

    Synthetic on purpose: producing a real .doc needs a writer library nobody
    should add for a signature test, and the detector's contract is exactly
    "OLE2 container whose directory names a WordDocument stream".
    """
    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    return header + "WordDocument".encode("utf-16-le") + b"\x00" * 64


ATTACHMENT_SAMPLES: dict[str, bytes] = {
    "png": make_png(),
    "jpeg": make_jpeg(),
    "pdf": make_pdf(),
    "zip": make_zip(),
    "gzip": make_gzip(),
    "rar": make_rar(),
    "doc": make_doc(),
    "docx": make_docx(),
    "xls": make_xls_sample(),
    "xlsx": make_xlsx_sample(),
}

# Filename each sample should be uploaded under.
ATTACHMENT_FILENAMES: dict[str, str] = {
    "png": "схема.png",
    "jpeg": "фото.jpg",
    "pdf": "договор.pdf",
    "zip": "пакет.zip",
    "gzip": "дамп.gz",
    "rar": "архив.rar",
    "doc": "письмо.doc",
    "docx": "письмо.docx",
    "xls": "смета.xls",
    "xlsx": "смета.xlsx",
}
