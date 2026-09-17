"""Full import cycle against the database (SPEC §6, §11)."""

from __future__ import annotations

import datetime as dt

import pytest
import sqlalchemy as sa

from app.core.access import AccessScope
from app.core.errors import ImportJobWrongStateError, InternalError
from app.imports.importer import ImportService
from app.models.enums import AuditAction, ImportJobStatus, ImportRowStatus, ImportTarget, UserRole
from app.models.interaction import Interaction
from app.models.university import University
from app.schemas.university import UniversityCreate
from app.schemas.user import UserCreate
from app.services.audit import search_audit_log
from app.services.catalogs import UniversityService
from app.services.interactions import InteractionService
from app.services.users import UserService
from tests.factories import catalog_row, make_xlsx


async def run_import(session, content: bytes, *, target=ImportTarget.INTERACTIONS):
    """Upload → auto-mapping → validate → commit, the happy path."""
    service = ImportService(session, AccessScope.system())
    created = await service.create_job(filename="catalog.xlsx", content=content, target=target)
    job = created["job"]
    await service.validate(job.id)
    _, stats = await service.commit(job.id)
    return service, job, stats


async def test_full_cycle_creates_every_entity(session):
    content = make_xlsx(
        [
            catalog_row(
                "МГТУ им. Баумана",
                vendor="Астра",
                product="Astra Linux",
                contract="ДЛ-001",
                signed="17.02.2026",
                years=3,
                contacts="Соколова Анна; Орлов Дмитрий",
                comment="первый импорт",
            )
        ]
    )
    service, job, stats = await run_import(session, content)

    assert stats["created"] == 1
    refreshed = await service.get_job(job.id)
    assert refreshed.status == ImportJobStatus.COMMITTED
    assert refreshed.committed_at is not None

    interaction = await session.scalar(sa.select(Interaction))
    assert interaction is not None
    assert interaction.contract_number == "ДЛ-001"
    assert interaction.license_signed_at == dt.date(2026, 2, 17)
    assert interaction.license_years == 3
    # Expiry is derived from signing date + term.
    assert interaction.license_expires_at == dt.date(2029, 2, 17)

    university = await session.scalar(sa.select(University))
    assert university is not None and university.name == "МГТУ им. Баумана"

    contacts = (await session.scalars(sa.text("SELECT full_name FROM university_contacts"))).all()
    assert set(contacts) == {"Соколова Анна", "Орлов Дмитрий"}


async def test_repeated_import_creates_no_duplicates(session):
    """SPEC §11: re-uploading the same file must be idempotent."""
    content = make_xlsx(
        [
            catalog_row("МГТУ им. Баумана", product="Astra Linux"),
            catalog_row("СПбПУ", product="РЕД ОС", vendor="Ред Софт"),
        ]
    )
    _, _, first = await run_import(session, content)
    assert first["created"] == 2

    _, _, second = await run_import(session, content)
    assert second["created"] == 0
    assert second["updated"] == 2

    for table in ("universities", "vendors", "it_products", "interactions"):
        count = await session.scalar(sa.text(f"SELECT count(*) FROM {table}"))
        assert count == 2, f"{table} got duplicated"

    # The second upload is flagged as a repeat, but is never blocked.
    created = await ImportService(session, AccessScope.system()).create_job(
        filename="catalog.xlsx", content=content, target=ImportTarget.INTERACTIONS
    )
    assert created["duplicate_of"] is not None
    assert created["warnings"]


async def test_dry_run_caches_repeated_university_and_manager_lookups(session, monkeypatch):
    university = await UniversityService(session, AccessScope.system()).create(
        UniversityCreate(name="MGTU")
    )
    manager = await UserService(session, AccessScope.system()).create(
        UserCreate(
            keycloak_id="kam-cache",
            full_name="Ivanov Ivan Ivanovich",
            role=UserRole.USER,
        )
    )
    content = make_xlsx([catalog_row(university.name, manager=manager.full_name) for _ in range(3)])
    service = ImportService(session, AccessScope.system())
    created = await service.create_job(
        filename="repeated.xlsx", content=content, target=ImportTarget.INTERACTIONS
    )

    university_calls = 0
    manager_calls = 0
    original_university_lookup = service.universities.repo.find_by_normalized_name
    original_manager_lookup = service.users.match_by_full_name

    async def count_university_lookup(normalized: str):
        nonlocal university_calls
        university_calls += 1
        return await original_university_lookup(normalized)

    async def count_manager_lookup(full_name: str):
        nonlocal manager_calls
        manager_calls += 1
        return await original_manager_lookup(full_name)

    monkeypatch.setattr(
        service.universities.repo, "find_by_normalized_name", count_university_lookup
    )
    monkeypatch.setattr(service.users, "match_by_full_name", count_manager_lookup)

    await service.validate(created["job"].id)

    assert university_calls == 1
    assert manager_calls == 1


async def test_empty_cell_does_not_erase_stored_value(session):
    """SPEC §11: an empty cell means "no data", never "clear the field"."""
    first = make_xlsx([catalog_row("МГТУ", contract="ДЛ-777", comment="важное примечание")])
    await run_import(session, first)

    second = make_xlsx([catalog_row("МГТУ", contract=None, comment=None)])
    await run_import(session, second)

    interaction = await session.scalar(sa.select(Interaction))
    assert interaction.contract_number == "ДЛ-777"
    assert interaction.comment == "важное примечание"


async def test_row_with_error_is_skipped_and_others_are_imported(session):
    """SPEC §11: one broken row must not hold the rest of the file hostage."""
    content = make_xlsx(
        [
            catalog_row("Первый вуз", product="ПО 1"),
            catalog_row("", product="ПО 2"),  # no university name -> error
            catalog_row("Третий вуз", product="ПО 3"),
        ]
    )
    service, job, stats = await run_import(session, content)

    assert stats["errors"] == 1
    assert stats["created"] == 2

    universities = await session.scalar(sa.text("SELECT count(*) FROM universities"))
    assert universities == 2

    rows, _ = await service.list_rows(job.id, status=ImportRowStatus.ERROR)
    assert len(rows) == 1
    assert rows[0].row_number == 3
    assert "Название ВУЗа" in rows[0].messages[0]["text"]


async def test_similar_university_is_flagged_not_merged(session):
    """SPEC §11: a near-match becomes a warning with a suggestion, never a merge."""
    await UniversityService(session, AccessScope.system()).create(
        UniversityCreate(name="Казанский федеральный университет")
    )
    content = make_xlsx([catalog_row("Казанский федеральный универститет")])
    service, job, _ = await run_import(session, content)

    rows, _ = await service.list_rows(job.id)
    assert rows[0].status == ImportRowStatus.WARNING
    suggestion = rows[0].messages[0]["suggestion"]
    assert suggestion["university_name"] == "Казанский федеральный университет"

    # Two separate universities exist: merging is a human decision.
    names = set((await session.scalars(sa.select(University.name))).all())
    assert names == {
        "Казанский федеральный университет",
        "Казанский федеральный универститет",
    }


async def test_duplicate_rows_inside_file_keep_the_last_one(session):
    content = make_xlsx(
        [
            catalog_row("МГТУ", product="Astra Linux", contract="СТАРЫЙ"),
            catalog_row("МГТУ", product="Astra Linux", contract="НОВЫЙ"),
        ]
    )
    _, _, stats = await run_import(session, content)

    assert stats["skipped"] == 1
    interaction = await session.scalar(sa.select(Interaction))
    assert interaction.contract_number == "НОВЫЙ"

    count = await session.scalar(sa.text("SELECT count(*) FROM interactions"))
    assert count == 1


async def test_unknown_manager_is_a_warning_and_row_still_imports(session):
    content = make_xlsx([catalog_row("МГТУ", manager="Неизвестный Сотрудник")])
    service, job, stats = await run_import(session, content)

    rows, _ = await service.list_rows(job.id)
    assert rows[0].status == ImportRowStatus.WARNING
    assert stats["created"] == 1

    interaction = await session.scalar(sa.select(Interaction))
    assert interaction.responsible_user_id is None


async def test_known_manager_is_matched_by_full_name(session, kam_user):
    content = make_xlsx([catalog_row("МГТУ", manager="  камов   кирилл андреевич ")])
    service, job, _ = await run_import(session, content)

    rows, _ = await service.list_rows(job.id)
    assert rows[0].status == ImportRowStatus.OK

    interaction = await session.scalar(sa.select(Interaction))
    assert interaction.responsible_user_id == kam_user.id


async def test_commit_is_all_or_nothing(session, monkeypatch):
    """SPEC §11: an artificial failure mid-commit must leave the DB untouched."""
    content = make_xlsx(
        [
            catalog_row("Первый вуз", product="ПО 1"),
            catalog_row("Второй вуз", product="ПО 2"),
            catalog_row("Третий вуз", product="ПО 3"),
        ]
    )
    service = ImportService(session, AccessScope.system())
    created = await service.create_job(
        filename="catalog.xlsx", content=content, target=ImportTarget.INTERACTIONS
    )
    job = created["job"]
    await service.validate(job.id)

    original = InteractionService.upsert
    calls = {"n": 0}

    async def exploding_upsert(self, **kwargs):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("искусственный сбой в середине импорта")
        return await original(self, **kwargs)

    monkeypatch.setattr(InteractionService, "upsert", exploding_upsert)

    with pytest.raises(InternalError):
        await service.commit(job.id)

    for table in ("universities", "vendors", "it_products", "interactions"):
        count = await session.scalar(sa.text(f"SELECT count(*) FROM {table}"))
        assert count == 0, f"{table} kept rows after a failed commit"

    failed = await service.get_job(job.id)
    assert failed.status == ImportJobStatus.FAILED
    assert failed.error_message


async def test_commit_writes_one_import_audit_entry(session):
    content = make_xlsx([catalog_row("МГТУ")])
    _, job, _ = await run_import(session, content)

    entries, total = await search_audit_log(session, action=AuditAction.IMPORT)
    assert total == 1
    assert entries[0].entity_id == job.id
    assert entries[0].changes["stats"]["created"] == 1

    # The entities themselves are audited as ordinary creates.
    _, created_total = await search_audit_log(
        session, entity_type="universities", action=AuditAction.CREATE
    )
    assert created_total == 1


async def test_commit_requires_validation_first(session):
    content = make_xlsx([catalog_row("МГТУ")])
    service = ImportService(session, AccessScope.system())
    created = await service.create_job(
        filename="catalog.xlsx", content=content, target=ImportTarget.INTERACTIONS
    )
    with pytest.raises(ImportJobWrongStateError):
        await service.commit(created["job"].id)


async def test_changing_mapping_resets_validation(session):
    content = make_xlsx([catalog_row("МГТУ")])
    service = ImportService(session, AccessScope.system())
    created = await service.create_job(
        filename="catalog.xlsx", content=content, target=ImportTarget.INTERACTIONS
    )
    job = created["job"]
    await service.validate(job.id)

    updated = await service.set_mapping(
        job.id, {"Название ВУЗа": "universities.name"}, save_as_preset="Только вуз"
    )
    assert updated.status == ImportJobStatus.PENDING

    presets, total = await service.list_presets(target=ImportTarget.INTERACTIONS)
    assert total == 1 and presets[0].name == "Только вуз"


async def test_import_of_universities_target(session):
    content = make_xlsx(
        [["МГТУ им. Баумана", "Бауманка", "г. Москва", "7701002520"]],
        headers=["Название ВУЗа", "Сокращённое название", "Регион", "ИНН"],
    )
    _, _, stats = await run_import(session, content, target=ImportTarget.UNIVERSITIES)
    assert stats["created"] == 1

    university = await session.scalar(sa.select(University))
    assert university.short_name == "Бауманка"
    assert university.region == "г. Москва"
    assert university.inn == "7701002520"


async def test_import_of_contacts_target(session):
    content = make_xlsx(
        [["МГТУ", "Соколова Анна", "Проректор", "a@example.edu", "+7 495 000-00-00"]],
        headers=["Название ВУЗа", "ФИО", "Должность", "Email", "Телефон"],
    )
    _, _, stats = await run_import(session, content, target=ImportTarget.CONTACTS)
    assert stats["created"] == 1

    row = (
        await session.execute(sa.text("SELECT full_name, position, email FROM university_contacts"))
    ).first()
    assert row == ("Соколова Анна", "Проректор", "a@example.edu")


async def test_import_of_products_target(session):
    content = make_xlsx(
        [["Astra Linux", "Астра", "DevOps; Информационная безопасность", "Российская ОС"]],
        headers=["ПО", "Вендор", "ИТ-направление", "Описание"],
    )
    _, _, stats = await run_import(session, content, target=ImportTarget.IT_PRODUCTS)
    assert stats["created"] == 1

    directions = await session.scalar(sa.text("SELECT count(*) FROM it_directions"))
    assert directions == 2
    links = await session.scalar(sa.text("SELECT count(*) FROM it_products_directions"))
    assert links == 2


async def test_report_contains_every_row_with_a_verdict(session):
    from openpyxl import load_workbook

    content = make_xlsx([catalog_row("Первый вуз"), catalog_row("")])
    service, job, _ = await run_import(session, content)

    from app.imports.report import build_report

    rows = await service.all_rows(job.id)
    report = build_report(await service.get_job(job.id), rows)

    workbook = load_workbook(__import__("io").BytesIO(report))
    sheet = workbook.active
    assert sheet.max_row == 3  # header + 2 rows
    verdicts = [sheet.cell(row=i, column=sheet.max_column - 2).value for i in (2, 3)]
    assert "Импортируется" in verdicts[0]
    assert verdicts[1] == "Не импортируется"
