"""The import flow as the frontend will drive it, over HTTP."""

from __future__ import annotations

from openpyxl import load_workbook

from tests.conftest import auth
from tests.factories import catalog_row, make_xlsx

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload_payload(content: bytes, target: str = "interactions"):
    return {
        "files": {"file": ("catalog.xlsx", content, XLSX_TYPE)},
        "data": {"target": target},
    }


async def test_full_http_cycle(session, client, manager_user):
    content = make_xlsx(
        [
            catalog_row("МГТУ им. Баумана", product="Astra Linux", contacts="Соколова Анна"),
            catalog_row("СПбПУ", product="РЕД ОС", vendor="Ред Софт"),
        ]
    )

    # Step 1 — upload.
    uploaded = await client.post(
        "/api/v1/imports", **upload_payload(content), headers=auth(manager_user)
    )
    assert uploaded.status_code == 201
    body = uploaded.json()
    job_id = body["job"]["id"]
    assert body["headers"][0] == "Название ВУЗа"
    assert body["duplicate_of"] is None
    suggested = {item["column"]: item["field"] for item in body["suggested_mapping"]}
    assert suggested["Название ВУЗа"] == "universities.name"

    # Step 2 — confirm the mapping the server proposed.
    mapping = await client.post(
        f"/api/v1/imports/{job_id}/mapping",
        json={"mapping": suggested},
        headers=auth(manager_user),
    )
    assert mapping.status_code == 200
    assert mapping.json()["status"] == "pending"

    # Step 3 — dry run.
    validated = await client.post(f"/api/v1/imports/{job_id}/validate", headers=auth(manager_user))
    assert validated.status_code == 200
    assert validated.json()["job"]["status"] == "validated"
    assert validated.json()["stats"]["to_create"] == 2

    rows = await client.get(f"/api/v1/imports/{job_id}/rows", headers=auth(manager_user))
    assert rows.json()["total"] == 2

    # Step 4 — commit.
    committed = await client.post(f"/api/v1/imports/{job_id}/commit", headers=auth(manager_user))
    assert committed.status_code == 200
    assert committed.json()["stats"]["created"] == 2
    assert committed.json()["job"]["status"] == "committed"

    interactions = await client.get("/api/v1/interactions", headers=auth(manager_user))
    assert interactions.json()["total"] == 2

    # Step 5 — report.
    report = await client.get(f"/api/v1/imports/{job_id}/report", headers=auth(manager_user))
    assert report.status_code == 200
    assert report.headers["content-type"] == XLSX_TYPE
    workbook = load_workbook(__import__("io").BytesIO(report.content))
    assert workbook.active.max_row == 3

    history = await client.get("/api/v1/imports", headers=auth(manager_user))
    assert history.json()["total"] == 1


async def test_upload_rejects_non_excel(session, client, manager_user):
    response = await client.post(
        "/api/v1/imports",
        files={"file": ("catalog.xlsx", b"not a spreadsheet at all", XLSX_TYPE)},
        data={"target": "interactions"},
        headers=auth(manager_user),
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "IMPORT_INVALID_FORMAT"


async def test_upload_rejects_empty_file(session, client, manager_user):
    response = await client.post(
        "/api/v1/imports",
        files={"file": ("catalog.xlsx", b"", XLSX_TYPE)},
        data={"target": "interactions"},
        headers=auth(manager_user),
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "IMPORT_INVALID_FORMAT"


async def test_mapping_without_required_field_is_reported(session, client, manager_user):
    content = make_xlsx([catalog_row("МГТУ")])
    uploaded = await client.post(
        "/api/v1/imports", **upload_payload(content), headers=auth(manager_user)
    )
    job_id = uploaded.json()["job"]["id"]

    response = await client.post(
        f"/api/v1/imports/{job_id}/mapping",
        json={"mapping": {"Вендор": "vendors.name"}},
        headers=auth(manager_user),
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "IMPORT_MAPPING_INCOMPLETE"


async def test_commit_before_validation_is_rejected(session, client, manager_user):
    content = make_xlsx([catalog_row("МГТУ")])
    uploaded = await client.post(
        "/api/v1/imports", **upload_payload(content), headers=auth(manager_user)
    )
    job_id = uploaded.json()["job"]["id"]

    response = await client.post(f"/api/v1/imports/{job_id}/commit", headers=auth(manager_user))
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "IMPORT_JOB_WRONG_STATE"


async def test_import_requires_manager_role(session, client, kam_user):
    content = make_xlsx([catalog_row("МГТУ")])
    response = await client.post(
        "/api/v1/imports", **upload_payload(content), headers=auth(kam_user)
    )
    assert response.status_code == 403


async def test_preset_can_be_created_and_listed(session, client, manager_user):
    created = await client.post(
        "/api/v1/imports/presets",
        json={
            "name": "Каталог ПО",
            "target": "interactions",
            "mapping": {"Название ВУЗа": "universities.name"},
        },
        headers=auth(manager_user),
    )
    assert created.status_code == 201

    listed = await client.get(
        "/api/v1/imports/presets", params={"target": "interactions"}, headers=auth(manager_user)
    )
    assert listed.json()["total"] == 1


async def test_preset_with_unknown_field_is_rejected(session, client, manager_user):
    response = await client.post(
        "/api/v1/imports/presets",
        json={"name": "Плохой", "target": "interactions", "mapping": {"X": "nope.nope"}},
        headers=auth(manager_user),
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


async def test_legacy_xls_goes_through_the_whole_cycle(session, client, manager_user):
    """A real customer file may still be .xls — the full cycle must accept it."""
    from tests.factories import make_xls

    content = make_xls([catalog_row("Вуз из .xls", product="Старое ПО", contract="ДЛ-XLS")])
    uploaded = await client.post(
        "/api/v1/imports",
        files={"file": ("catalog.xls", content, "application/vnd.ms-excel")},
        data={"target": "interactions"},
        headers=auth(manager_user),
    )
    assert uploaded.status_code == 201
    job_id = uploaded.json()["job"]["id"]

    validated = await client.post(f"/api/v1/imports/{job_id}/validate", headers=auth(manager_user))
    assert validated.json()["stats"]["to_create"] == 1

    committed = await client.post(f"/api/v1/imports/{job_id}/commit", headers=auth(manager_user))
    assert committed.json()["stats"]["created"] == 1

    interactions = await client.get("/api/v1/interactions", headers=auth(manager_user))
    assert interactions.json()["items"][0]["contract_number"] == "ДЛ-XLS"
