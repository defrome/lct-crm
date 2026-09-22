"""Reading .xls/.xlsx into raw rows (SPEC §6 step 1).

The parser knows nothing about the domain: it produces headers and cell values
and leaves every interpretation to `validators.py`. That separation is what lets
SPEC-04's integrations reuse the validation rules without an Excel file.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from itertools import chain
from typing import Any, Literal

from openpyxl import load_workbook

from app.core.config import settings
from app.core.errors import ImportFileTooLargeError, ImportInvalidFormatError

FileFormat = Literal["xlsx", "xls"]

# Magic numbers. The client-supplied filename and Content-Type are not trusted
# (SPEC §6 step 1: "реальный MIME по сигнатуре файла").
_ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
_OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_REPORT_METADATA = re.compile(
    r"^Отбор:\s*.+?\.\s*Записей:\s*\d+\.\s*Сформирован\s+.+$",
    re.IGNORECASE,
)


@dataclass(slots=True)
class RawRow:
    """One spreadsheet row, keyed by column header."""

    row_number: int  # 1-based, as shown in Excel (header is row 1)
    cells: dict[str, Any] = field(default_factory=dict)

    def is_empty(self) -> bool:
        return all(value is None or str(value).strip() == "" for value in self.cells.values())


@dataclass(slots=True)
class ParsedFile:
    headers: list[str]
    rows: list[RawRow]
    file_hash: str
    file_format: FileFormat
    sheet_name: str


def _is_report_metadata_row(row: RawRow) -> bool:
    """Ignore the one-cell filter summary emitted by the report exporter."""
    return _is_report_metadata_values(row.cells.values())


def _is_report_metadata_values(values: Any) -> bool:
    """Recognise the filter summary that precedes headers in exported reports."""
    values = [value for value in values if value is not None and str(value).strip()]
    return (
        len(values) == 1
        and isinstance(values[0], str)
        and bool(_REPORT_METADATA.fullmatch(values[0]))
    )


def compute_file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def detect_format(content: bytes, filename: str) -> FileFormat:
    """Determine the real format from the file's own bytes."""
    if content.startswith(_OLE2_SIGNATURE):
        return "xls"
    if content.startswith(_ZIP_SIGNATURES):
        # A .docx is also a zip; require the OOXML spreadsheet marker.
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            raise ImportInvalidFormatError(
                "Файл повреждён и не может быть прочитан как таблица Excel",
                details={"filename": filename},
            ) from None
        if "xl/workbook.xml" in names:
            return "xlsx"
    raise ImportInvalidFormatError(
        "Файл не является корректной таблицей Excel",
        details={"filename": filename},
    )


def _normalize_header(value: Any, index: int) -> str:
    """Headers are the mapping keys, so they must be unique and non-empty."""
    text = "" if value is None else str(value).strip()
    text = " ".join(text.split())
    return text or f"Колонка {index + 1}"


def _cell_value(value: Any) -> Any:
    """Normalise a cell into something JSON-serialisable.

    Dates keep their type information through an ISO string; the validator
    re-parses it. Trailing/leading whitespace is stripped here once so every
    later stage sees clean text.
    """
    if value is None:
        return None
    if isinstance(value, str):
        text = " ".join(value.split())
        return text or None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _parse_xlsx(content: bytes, filename: str) -> tuple[list[str], list[RawRow], str]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise ImportInvalidFormatError(
            "Не удалось прочитать файл Excel",
            details={"filename": filename, "reason": str(exc)},
        ) from None

    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            raise ImportInvalidFormatError(
                "В файле нет ни одной строки", details={"filename": filename}
            ) from None

        # Browser-generated report files contain a title and a one-cell filter
        # summary before their headers.  Treat the row after that summary as
        # the header so an exported, selected set of columns can be imported
        # again without manual mapping.
        try:
            next_row = next(iterator)
        except StopIteration:
            next_row = None
        if next_row is not None and _is_report_metadata_values(next_row):
            try:
                header_row = next(iterator)
            except StopIteration:
                raise ImportInvalidFormatError(
                    "В отчёте отсутствует строка заголовков", details={"filename": filename}
                ) from None
            start_row = 4
        else:
            start_row = 2

        headers = _dedupe([_normalize_header(value, i) for i, value in enumerate(header_row)])
        rows: list[RawRow] = []
        if next_row is not None and start_row == 2:
            iterator = chain((next_row,), iterator)
        for offset, values in enumerate(iterator, start=start_row):
            cells = {
                header: _cell_value(values[i]) if i < len(values) else None
                for i, header in enumerate(headers)
            }
            row = RawRow(row_number=offset, cells=cells)
            if not row.is_empty() and not _is_report_metadata_row(row):
                rows.append(row)
        return headers, rows, sheet.title
    finally:
        workbook.close()


def _parse_xls(content: bytes, filename: str) -> tuple[list[str], list[RawRow], str]:
    # Imported lazily: legacy .xls is the rare case and xlrd should not be a
    # hard import-time dependency of the whole application.
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise ImportInvalidFormatError(
            "Не удалось прочитать файл Excel (формат .xls)",
            details={"filename": filename, "reason": str(exc)},
        ) from None

    sheet = book.sheet_by_index(0)
    if sheet.nrows == 0:
        raise ImportInvalidFormatError(
            "В файле нет ни одной строки", details={"filename": filename}
        )

    header_index = 0
    if sheet.nrows > 1 and _is_report_metadata_values(sheet.row_values(1)):
        if sheet.nrows <= 2:
            raise ImportInvalidFormatError(
                "В отчёте отсутствует строка заголовков", details={"filename": filename}
            )
        header_index = 2

    headers = _dedupe(
        [_normalize_header(sheet.cell_value(header_index, col), col) for col in range(sheet.ncols)]
    )
    rows: list[RawRow] = []
    for row_index in range(header_index + 1, sheet.nrows):
        cells: dict[str, Any] = {}
        for col, header in enumerate(headers):
            cell = sheet.cell(row_index, col)
            value: Any = cell.value
            # xlrd hands dates back as floats plus a type flag; convert them
            # here so downstream code never sees an Excel serial from .xls.
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = dt.datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_EMPTY:
                value = None
            cells[header] = _cell_value(value)
        row = RawRow(row_number=row_index + 1, cells=cells)
        if not row.is_empty() and not _is_report_metadata_row(row):
            rows.append(row)
    return headers, rows, sheet.name


def _dedupe(headers: list[str]) -> list[str]:
    """Excel happily allows two columns with the same title; the mapping cannot."""
    seen: dict[str, int] = {}
    result: list[str] = []
    for header in headers:
        if header in seen:
            seen[header] += 1
            result.append(f"{header} ({seen[header]})")
        else:
            seen[header] = 1
            result.append(header)
    return result


def parse_file(content: bytes, filename: str) -> ParsedFile:
    """Full step-1 check: size, real format, at least one data row."""
    if len(content) > settings.import_max_file_size:
        raise ImportFileTooLargeError(
            "Размер файла превышает допустимый предел",
            details={
                "filename": filename,
                "size": len(content),
                "max_size": settings.import_max_file_size,
            },
        )
    if not content:
        raise ImportInvalidFormatError("Файл пуст", details={"filename": filename})

    file_format = detect_format(content, filename)
    if file_format == "xlsx":
        headers, rows, sheet_name = _parse_xlsx(content, filename)
    else:
        headers, rows, sheet_name = _parse_xls(content, filename)

    if not rows:
        raise ImportInvalidFormatError(
            "В файле нет ни одной непустой строки данных",
            details={"filename": filename},
        )
    if len(rows) > settings.import_max_rows:
        raise ImportFileTooLargeError(
            "В файле слишком много строк",
            details={"rows": len(rows), "max_rows": settings.import_max_rows},
        )

    return ParsedFile(
        headers=headers,
        rows=rows,
        file_hash=compute_file_hash(content),
        file_format=file_format,
        sheet_name=sheet_name,
    )
