"""Error report as .xlsx (SPEC §6 step 5).

The report keeps the original columns untouched and appends the verdict, so the
user can fix the cells in place and re-upload the same file.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.enums import ImportRowStatus
from app.models.import_job import ImportJob, ImportRow

_STATUS_TITLES = {
    ImportRowStatus.OK.value: "Импортируется",
    ImportRowStatus.WARNING.value: "Импортируется с замечаниями",
    ImportRowStatus.ERROR.value: "Не импортируется",
}

_STATUS_FILLS = {
    ImportRowStatus.OK.value: PatternFill("solid", fgColor="E7F4E4"),
    ImportRowStatus.WARNING.value: PatternFill("solid", fgColor="FFF4CE"),
    ImportRowStatus.ERROR.value: PatternFill("solid", fgColor="FBE3E4"),
}

RESULT_COLUMNS = ("Результат импорта", "Замечания", "№ строки в файле")


def _message_text(message: dict[str, Any]) -> str:
    text = str(message.get("text", "")).strip()
    field = message.get("field")
    return f"[{field}] {text}" if field else text


def build_report(job: ImportJob, rows: list[ImportRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат импорта"

    headers = [str(header) for header in job.source_headers]
    sheet.append([*headers, *RESULT_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in sorted(rows, key=lambda item: item.row_number):
        status = row.status.value if hasattr(row.status, "value") else str(row.status)
        messages = "; ".join(_message_text(m) for m in (row.messages or []))
        if row.parsed_data.get("superseded_by_row"):
            messages = (messages + "; " if messages else "") + "строка пропущена как дубль"
        values = [row.raw_data.get(header) for header in headers]
        sheet.append([*values, _STATUS_TITLES.get(status, status), messages, row.row_number])

        fill = _STATUS_FILLS.get(status)
        if fill is not None:
            sheet.cell(row=sheet.max_row, column=len(headers) + 1).fill = fill

    # Readable widths: enough for the content, capped so one long comment does
    # not make the sheet unusable.
    for index in range(1, len(headers) + len(RESULT_COLUMNS) + 1):
        letter = get_column_letter(index)
        longest = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 60)

    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
