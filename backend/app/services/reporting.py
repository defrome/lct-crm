"""Reporting queries and file renderers.

The database query is deliberately shared by tabular reports and charts.  This
keeps the numbers in a chart equal to the numbers a user can export with the
same filters.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import os
from collections import Counter
from typing import Any

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.access import AccessScope
from app.core.errors import ValidationError
from app.models.interaction import Interaction
from app.models.product import ITDirection, ITProduct
from app.models.university import University
from app.models.user import User
from app.models.workflow import WorkflowStage
from app.repositories.base import visible_university_ids
from app.schemas.report import ReportColumn, ReportFilters, ReportFormat

COLUMN_TITLES: dict[ReportColumn, str] = {
    ReportColumn.UNIVERSITY: "University",
    ReportColumn.IT_DIRECTION: "IT direction",
    ReportColumn.IT_PRODUCT: "IT product",
    ReportColumn.WORKFLOW_STAGE: "Workflow stage",
    ReportColumn.RESPONSIBLE: "Responsible",
}

_COLUMN_EXPRESSIONS: dict[ReportColumn, Any] = {
    ReportColumn.UNIVERSITY: University.name,
    ReportColumn.IT_DIRECTION: ITDirection.name,
    ReportColumn.IT_PRODUCT: ITProduct.name,
    ReportColumn.WORKFLOW_STAGE: WorkflowStage.name,
    ReportColumn.RESPONSIBLE: User.full_name,
}


def period_conditions(period_from: dt.date | None, period_to: dt.date | None) -> list[Any]:
    """Return an inclusive validity-period filter for a licence."""
    if period_from and period_to and period_from > period_to:
        raise ValidationError("Начало периода не может быть позже конца")
    if period_from and period_to:
        return [
            sa.or_(
                Interaction.license_signed_at.between(period_from, period_to),
                Interaction.license_expires_at.between(period_from, period_to),
                sa.and_(
                    Interaction.license_signed_at <= period_from,
                    Interaction.license_expires_at >= period_to,
                ),
            )
        ]
    if period_from:
        return [
            sa.or_(
                Interaction.license_signed_at >= period_from,
                Interaction.license_expires_at >= period_from,
            )
        ]
    if period_to:
        return [
            sa.or_(
                Interaction.license_signed_at <= period_to,
                Interaction.license_expires_at <= period_to,
            )
        ]
    return []


class ReportingService:
    def __init__(self, session: AsyncSession, scope: AccessScope | None = None) -> None:
        self.session = session
        self.scope = scope or AccessScope.system()

    def _statement(self, filters: ReportFilters, columns: list[ReportColumn]) -> Any:
        selected = [_COLUMN_EXPRESSIONS[column].label(column.value) for column in columns]
        stmt = (
            sa.select(*selected)
            .select_from(Interaction)
            .outerjoin(University, Interaction.university_id == University.id)
            .outerjoin(ITDirection, Interaction.it_direction_id == ITDirection.id)
            .outerjoin(ITProduct, Interaction.it_product_id == ITProduct.id)
            .outerjoin(WorkflowStage, Interaction.current_stage_id == WorkflowStage.id)
            .outerjoin(User, Interaction.responsible_user_id == User.id)
            .where(Interaction.deleted_at.is_(None))
        )
        conditions: list[Any] = []
        for field in (
            "university_id",
            "it_direction_id",
            "it_product_id",
            "responsible_user_id",
        ):
            value = getattr(filters, field)
            if value is not None:
                conditions.append(getattr(Interaction, field) == value)
        conditions.extend(period_conditions(filters.period_from, filters.period_to))
        if not self.scope.is_privileged:
            conditions.append(Interaction.university_id.in_(visible_university_ids(self.scope)))
        return stmt.where(*conditions)

    async def rows(
        self, filters: ReportFilters, columns: list[ReportColumn]
    ) -> list[dict[str, str]]:
        result = await self.session.execute(
            self._statement(filters, columns).order_by(University.name, Interaction.id)
        )
        return [
            {column.value: str(row[column.value] or "") for column in columns}
            for row in result.mappings().all()
        ]


def _workbook(rows: list[dict[str, str]], columns: list[ReportColumn]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([COLUMN_TITLES[column] for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row[column.value] for column in columns])
    for index in range(1, len(columns) + 1):
        letter = get_column_letter(index)
        longest = max((len(str(cell.value or "")) for cell in sheet[letter]), default=10)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 60)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _xls(rows: list[dict[str, str]], columns: list[ReportColumn]) -> bytes:
    import xlwt

    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Report")
    header_style = xlwt.easyxf("font: bold on")
    for index, column in enumerate(columns):
        sheet.write(0, index, COLUMN_TITLES[column], header_style)
    for row_index, row in enumerate(rows, start=1):
        for column_index, column in enumerate(columns):
            sheet.write(row_index, column_index, row[column.value])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _font_path() -> str | None:
    candidates = (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
    )
    return next((path for path in candidates if os.path.exists(path)), None)


def _pdf(rows: list[dict[str, str]], columns: list[ReportColumn]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    font_name = "Helvetica"
    path = _font_path()
    if path:
        font_name = "ReportFont"
        pdfmetrics.registerFont(TTFont(font_name, path))
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    style = getSampleStyleSheet()["Heading2"]
    style.fontName = font_name
    data = [[COLUMN_TITLES[column] for column in columns]]
    data.extend([[row[column.value] for column in columns] for row in rows])
    widths = [(landscape(A4)[0] - 24 * mm) / len(columns)] * len(columns)
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    document.build([Paragraph("CRM interactions report", style), Spacer(1, 5 * mm), table])
    return buffer.getvalue()


def render_report(
    rows: list[dict[str, str]], columns: list[ReportColumn], format_: ReportFormat
) -> bytes:
    if format_ == ReportFormat.XLSX:
        return _workbook(rows, columns)
    if format_ == ReportFormat.XLS:
        return _xls(rows, columns)
    if format_ == ReportFormat.PDF:
        return _pdf(rows, columns)
    return json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")


def render_chart(groups: Counter[str], title: str, *, as_pdf: bool) -> bytes:
    """Render a deterministic bar chart from the same grouped report data."""
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1200, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = _font_path()
    title_font = ImageFont.truetype(font_path, 28) if font_path else ImageFont.load_default()
    label_font = ImageFont.truetype(font_path, 16) if font_path else ImageFont.load_default()
    draw.text((50, 30), title, fill="#172B4D", font=title_font)
    items = sorted(groups.items(), key=lambda item: (-item[1], item[0]))
    if len(items) > 20:
        shown, remaining = items[:19], items[19:]
        items = [*shown, ("Other", sum(value for _, value in remaining))]
    if not items:
        items = [("No data", 0)]
    max_value = max(value for _, value in items) or 1
    chart_left, chart_top, chart_bottom = 120, 110, 620
    chart_width = width - chart_left - 60
    gap = 16
    bar_width = max(28, (chart_width - gap * (len(items) - 1)) // len(items))
    for index, (label, value) in enumerate(items):
        left = chart_left + index * (bar_width + gap)
        bar_height = int((chart_bottom - chart_top) * value / max_value)
        draw.rectangle(
            (left, chart_bottom - bar_height, left + bar_width, chart_bottom), fill="#247BA0"
        )
        draw.text((left, chart_bottom + 8), label[:18], fill="#172B4D", font=label_font)
        draw.text(
            (left, chart_bottom - bar_height - 24), str(value), fill="#172B4D", font=label_font
        )
    draw.line(
        (chart_left, chart_top, chart_left, chart_bottom, width - 60, chart_bottom),
        fill="#52606D",
        width=2,
    )
    buffer = io.BytesIO()
    image.save(buffer, format="PDF" if as_pdf else "PNG", resolution=144.0)
    return buffer.getvalue()
